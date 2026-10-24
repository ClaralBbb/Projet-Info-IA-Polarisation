"""Prépare un pilote d'annotation humain, sans appel LLM ni API."""
import argparse
import html
import json
import random
from collections import Counter, defaultdict
from pathlib import Path
from urllib.parse import quote
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from analyze import load, tokens, duplicate_key, THEMES

SEED=20261024
LEXICON=THEMES['Police et maintien de l’ordre'] | THEMES['Violences et dégâts'] | set('flic flics matraque matraques lbd flashball repression repressif repressifs brutalite brutalites bavure bavures tirs tir gaz affrontements tabasse tabassent tabasser violent violente violents violentes'.split())

def select(comments):
    rng=random.Random(SEED)
    order=list(comments);rng.shuffle(order)
    seen=set();buckets=defaultdict(list)
    for row in order:
        key=duplicate_key(str(row.get('text') or ''))
        if key in seen: continue
        if not set(tokens(str(row.get('text') or ''))) & LEXICON: continue
        seen.add(key)
        buckets[(bool(row.get('parent_id')),row['video_id'])].append(row)
    selected=[];used=Counter()
    for replies in [False,True]:
        vids=[vid for kind,vid in buckets if kind==replies]
        rng.shuffle(vids)
        n=0
        for _ in range(2):
            for vid in vids:
                bucket=buckets[(replies,vid)]
                if n<50 and used[vid]<2 and bucket:
                    selected.append(bucket.pop());used[vid]+=1;n+=1
        if n!=50: raise ValueError('Pas assez de candidats pour les quotas 50/50 et deux par vidéo maximum')
    rng.shuffle(selected)
    return selected

def sheet(wb,name,headers,rows,widths=None):
    ws=wb.create_sheet(name);ws.append(headers)
    for row in rows: ws.append(row)
    for row in ws:
        for cell in row:
            if isinstance(cell.value,str): cell.data_type='s'
            cell.alignment=Alignment(vertical='top',wrap_text=True)
    for cell in ws[1]:
        cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='244766')
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for i in range(1,len(headers)+1):
        ws.column_dimensions[ws.cell(1,i).column_letter].width=(widths or {}).get(i,25)
    return ws

def dropdown(ws,column,options,last=1001):
    rule=DataValidation(type='list',formula1='"'+','.join(options)+'"',allow_blank=True)
    rule.errorTitle='Valeur non prévue';rule.error='Choisir une valeur de la liste.'
    rule.showErrorMessage=True;ws.add_data_validation(rule);rule.add(f'{column}2:{column}{last}')

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--exclude',type=Path)
    parser.add_argument('--output',type=Path,default=Path('annotation'))
    args=parser.parse_args()
    source=next(p for p in [Path('data/paris_2026-10-17.json'),Path('data/paris_2026-10-17.xlsx')] if p.exists())
    videos,comments=load(source)
    exclusion=json.loads(args.exclude.read_text()) if args.exclude else {}
    excluded=set(exclusion.get('excluded_comment_ids',[]))
    excluded_videos=set(exclusion.get('excluded_video_ids',[]))
    selected=select([c for c in comments if c['comment_id'] not in excluded and c['video_id'] not in excluded_videos])
    video_map={v['video_id']:v for v in videos};lookup={c['comment_id']:c for c in comments}
    wb=Workbook();wb.remove(wb.active)
    guide=[
        ['Objectif','Pilote humain : préciser la grille police/violences avant toute classification IA. Aucun appel API.'],
        ['Étape 1','Lire Commentaires, texte complet et parent référencé. Vérifier le lien avec le sujet et renseigner pertinence_sujet. Ne pas supposer que la vidéo est validée.'],
        ['Étape 2','Dans Positions, une ligne est préparée par commentaire. Identifier acteur, objet, position et citer un passage exact. Pour un commentaire hors sujet, ne pas inventer de position.'],
        ['Plusieurs positions','Ajouter une ligne dans Positions avec le même comment_id et un autre position_id, par exemple C001-P2. Séparer les objets et positions distincts.'],
        ['Responsabilité','Renseigner seulement si le texte attribue explicitement à un acteur la responsabilité de quelque chose. Décrire cette attribution, pas un fait établi.'],
        ['Argument','Résumer la justification donnée dans le texte. Laisser vide si elle est absente. Ne pas inventer les intentions de l’auteur.'],
        ['Ambiguïté','Utiliser contexte insuffisant ou ambigu si ironie, citation, pronoms ou désaccord implicite empêchent une lecture fiable. Ne pas forcer une étiquette.'],
        ['Méthode de sélection','Graine 20261024. Présélection lexicale, textes normalisés distincts, 50 commentaires principaux et 50 réponses, maximum deux commentaires par vidéo. Échantillon de découverte, non représentatif.'],
        ['Contexte disponible','Titre et chaîne, commentaire complet et parent référencé (souvent le commentaire principal). Pas le fil complet, pas la vidéo visionnée ni sa transcription.'],
        ['Texte et vie privée','Texte conservé pour l’interprétation; peut contenir propos offensants ou données personnelles. Aucun nom de compte auteur ajouté.'],
        ['À faire maintenant','Annoter d’abord 15 à 20 commentaires. Noter les difficultés dans note. Discuter la grille avant de terminer les 100.'],
        ['Évaluation future','Ces 100 commentaires servent à construire la grille. Un nouvel échantillon séparé servira à comparer humain et IA sans montrer les réponses IA à l’annotatrice.'],
        ['Vocabulaire','acteur = qui est évalué; objet = action ou rôle évalué; position = attitude vis-à-vis de cet objet. Une critique de méthode ne signifie pas opposition générale à la police.'],
    ]
    sheet(wb,'Guide',['Rubrique','Consigne'],guide,{1:28,2:110})
    rows=[];annotations=[]
    for i,c in enumerate(selected,1):
        ref=f'C{i:03}';v=video_map.get(c['video_id'],{});parent=lookup.get(c.get('parent_id'))
        url='https://www.youtube.com/watch?v='+quote(c['video_id'])+'&lc='+quote(c['comment_id'])
        rows.append([ref,c['comment_id'],c['video_id'],html.unescape(v.get('title','')),html.unescape(v.get('channel','')),url,
            'réponse' if c.get('parent_id') else 'principal',c.get('text',''),c.get('parent_id',''),
            parent.get('text','') if parent else '', 'disponible' if parent else ('manquant' if c.get('parent_id') else 'sans parent'),
            ' | '.join(sorted(set(tokens(c.get('text',''))) & LEXICON)),'','',''])
        annotations.append([ref+'-P1',ref,c['comment_id'],'','','','','','','',''])
    ws=sheet(wb,'Commentaires',['repere','comment_id','video_id','titre_video','chaine','lien_commentaire','type','texte_complet','parent_id','texte_parent','statut_parent','mots_selection','pertinence_sujet','pertinence_video','note_contexte'],rows,{4:55,6:35,8:95,10:95,12:30,15:45})
    for row in range(2,102):
        ws.cell(row,6).hyperlink=ws.cell(row,6).value;ws.cell(row,6).style='Hyperlink';ws.row_dimensions[row].height=90
    dropdown(ws,'M',['oui','non','incertain']);dropdown(ws,'N',['oui','non','non vérifiée'])
    pos=sheet(wb,'Positions',['position_id','repere','comment_id','acteur','objet','position','responsabilite_attribuee','argument','preuve_textuelle','incertitude','note'],annotations,{5:40,7:45,8:55,9:70,11:45})
    dropdown(pos,'D',['police','manifestants','casseurs','gouvernement','autre','indéterminé'])
    dropdown(pos,'E',['usage de la force','interpellations','protection et sécurité','violences et dégâts','rôle politique','autre','indéterminé'])
    dropdown(pos,'F',['condamne','justifie ou soutient','nuance','décrit sans position','indéterminée'])
    dropdown(pos,'J',['lecture explicite','ambigu','contexte insuffisant'])
    examples=[
        ['FICTIF 1','La police frappe des gens pacifiques.','police','usage de la force','condamne','frappe des gens pacifiques','Critique des méthodes, pas nécessairement de toute la police.'],
        ['FICTIF 2','La police protège les passants contre les casseurs.','police','protection et sécurité','justifie ou soutient','protège les passants','La protection est présentée favorablement.'],
        ['FICTIF 3 — position 1','Arrêtez les casseurs, mais ne frappez pas les gens pacifiques.','police','interpellations','justifie ou soutient','Arrêtez les casseurs','Deux positions dans un commentaire : créer deux lignes.'],
        ['FICTIF 3 — position 2','Arrêtez les casseurs, mais ne frappez pas les gens pacifiques.','police','usage de la force','condamne','ne frappez pas les gens pacifiques','Même commentaire que la ligne précédente.'],
        ['FICTIF 4','Ils disent que la police est trop violente, je ne sais pas.','police','usage de la force','indéterminée','je ne sais pas','Ne pas attribuer à l’auteur la position qu’il cite.'],
        ['FICTIF 5','La police protège la droite.','police','rôle politique','indéterminée','protège la droite','Attribution de rôle politique explicite; approbation ou condamnation non explicite.'],
    ]
    sheet(wb,'Exemples fictifs',['exemple','texte_inventé','acteur','objet','position','preuve','explication'],examples,{2:70,6:50,7:75})
    out=args.output;out.mkdir(parents=True,exist_ok=True)
    path=out/'pilote_police_violences_100.xlsx';wb.save(path)
    check=load_workbook(path,read_only=True)
    assert check['Commentaires'].max_row==101
    assert check['Positions'].max_row==101
    assert all(all(v is None for v in row[3:]) for row in check['Positions'].iter_rows(min_row=2,values_only=True))
    check.close()
    summary={'seed':SEED,'source':source.name,'comments':100,'videos':len({r['video_id'] for r in selected}),
        'types':dict(Counter('réponse' if r.get('parent_id') else 'principal' for r in selected)),
        'parents_missing':sum(bool(r.get('parent_id')) and r['parent_id'] not in lookup for r in selected),
        'lexicon':sorted(LEXICON),'selected_comment_ids':[r['comment_id'] for r in selected]}
    (out/'selection.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2))
    print(json.dumps({k:v for k,v in summary.items() if k not in ['lexicon','selected_comment_ids']},ensure_ascii=False,indent=2))
    print('Excel:',path.stat().st_size,'octets')

if __name__=='__main__':main()
