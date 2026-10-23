"""Analyse lexicale descriptive, sans inférence d'opinions individuelles."""
import argparse
import csv
import hashlib
import html
import json
import random
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

STOP_TEXT = """a à au aux avec ce ces cet cette ceux chaque chez comme comment dans de des du elle elles en entre et eux fait faire faut ils il je la le les leur leurs lui ma mais me mes moi mon même ne nos notre nous on ont ou où par pas plus pour pourquoi quand que quel quelle qui quoi sa sans se ses si son sont sous sur ta te tes toi ton tous tout toute toutes trop tu un une vos votre vous y ça est c d j l m n qu s t était être avoir ai as avons avez aussi bien bon donc encore été très va vais allez alors oui non rien dit dire ici là peu peut deux parce car merci vraiment déjà maintenant toujours jamais après avant seulement autres autre pendant contre ainsi dont quels quelles soit suis es the and of to in is it for you that this are be not with have they from was we an or but at your"""

def normalize(text):
    return ''.join(c for c in unicodedata.normalize('NFKD', text.lower()) if not unicodedata.combining(c))

STOP = set(normalize(STOP_TEXT).split()) | {'aujourd','hui','depuis'} | set('what who how why do does did will would can could should there their them he she his her its our us were been being these those than then so such if when which where all any some more most much many into out up down just very no yes because'.split())
THEMES = {
    'Police et maintien de l’ordre': {'police','policiers','policier','crs','gendarme','gendarmes','lacrymogene','lacrymogenes','interpellation','interpellations'},
    'Violences et dégâts': {'violence','violences','casseurs','casseur','pillage','pillages','casse','emeute','emeutes','incendie','incendies'},
    'Mobilisation et revendications': {'manifestation','manifestations','manifestants','manifestant','greve','syndicat','syndicats','mobilisation','revendications'},
    'Institutions et responsables politiques': {'gouvernement','president','ministre','macron','democratie','republique','etat','politique','politiques'},
    'Conditions de vie et économie': {'salaire','salaires','retraite','retraites','inflation','pauvrete','prix','impots','chomage','pouvoir'},
    'Médias et information': {'medias','media','journaliste','journalistes','propagande','information','desinformation','television','bfm','blast','mediapart'},
}
TOKEN = re.compile(r'[a-z]+')
URL = re.compile(r'https?://\S+|www\.\S+', re.I)

def tokens(text):
    return [x for x in TOKEN.findall(normalize(URL.sub(' ', html.unescape(text)))) if len(x)>2 and x not in STOP]

def phrases(text):
    # Conserver l'adjacence originale : ne pas fabriquer de paires à travers les mots vides.
    result = []
    for segment in re.split(r'[.!?;:\n\d]+', URL.sub('.', html.unescape(text))):
        words = TOKEN.findall(normalize(segment))
        result.extend(a+' '+b for a,b in zip(words,words[1:]) if len(a)>2 and len(b)>2 and a not in STOP and b not in STOP)
    return set(result)

def duplicate_key(text):
    return re.sub(r'\s+',' ',unicodedata.normalize('NFKC',html.unescape(text)).casefold()).strip()

def redact(text):
    text=URL.sub('[lien]',text)
    text=re.sub(r'\b[^\s@]+@[^\s@]+\.[^\s@]+\b','[email]',text)
    text=re.sub(r'(?<!\w)@[^\s]+','[mention]',text)
    text=re.sub(r'(?<!\w)(?:\+?\d[\s().-]*){9,15}(?!\w)','[numéro]',text)
    return text

def load(path):
    if path.suffix=='.json':
        data=json.loads(path.read_text())
        return data['videos'],data['comments']
    from openpyxl import load_workbook
    wb=load_workbook(path,read_only=True,data_only=True)
    def rows(name):
        iterator=wb[name].iter_rows(values_only=True)
        headers=next(iterator)
        return [dict(zip(headers,row)) for row in iterator]
    videos,comments=rows('Vidéos'),rows('Commentaires')
    wb.close()
    return videos,comments

def csv_write(path, headers, rows):
    with path.open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.writer(f)
        writer.writerow(headers)
        for row in rows:
            writer.writerow(["'"+x if isinstance(x,str) and x.startswith(('=','+','-','@','\r','\t')) else x for x in row])

def analyze(videos, comments):
    profiles={k:{'n':0,'words':Counter(),'pairs':Counter(),'themes':Counter()} for k in ['tous','uniques','principaux','reponses']}
    seen=set(); quality=Counter(); by_video=Counter(); theme_videos=defaultdict(set)
    sample_by_video={}; rng=random.Random(42); seen_in_video=Counter()
    for row in comments:
        text=str(row.get('text') or '').replace('\r\n','\n').replace('\r','\n')
        key=duplicate_key(text)
        ts=set(tokens(text))
        duplicate=key in seen
        if not key: quality['vides']+=1
        if duplicate: quality['doublons_textuels']+=1
        seen.add(key)
        if not ts:
            quality['sans_mot_analyse']+=1
            continue
        quality['analysables']+=1
        if len(ts)<3: quality['moins_de_3_mots_distincts']+=1
        vid=row['video_id']; by_video[vid]+=1
        found=[theme for theme,lexicon in THEMES.items() if ts & lexicon]
        ps=phrases(text)
        keys=['tous', 'reponses' if row.get('parent_id') else 'principaux']
        if not duplicate: keys.append('uniques')
        for profile in keys:
            p=profiles[profile];p['n']+=1;p['words'].update(ts);p['pairs'].update(ps);p['themes'].update(found)
        for theme in found: theme_videos[theme].add(vid)
        # Un commentaire par vidéo via réservoir; puis au plus 400 vidéos tirées au sort.
        seen_in_video[vid]+=1
        if rng.randrange(seen_in_video[vid])==0:
            sample_by_video[vid]={'video_id':vid,'comment_id':row['comment_id'],
                'text':redact(text)[:1200], 'themes':found,
                'type':'réponse' if row.get('parent_id') else 'principal'}
    chosen=rng.sample(sorted(sample_by_video),min(400,len(sample_by_video)))
    sample=[sample_by_video[vid] for vid in chosen]
    lookup={v['video_id']:v for v in videos}
    for r in sample:
        r['title']=lookup.get(r['video_id'],{}).get('title','')
        r['url']='https://www.youtube.com/watch?v='+r['video_id']
    result={'created_at':datetime.now(timezone.utc).isoformat(),'total':len(comments),
        'videos':len(videos),'videos_with_analyzable_comments':len(by_video),
        'quality':dict(quality), 'sample':sample,
        'theme_videos':{k:len(v) for k,v in theme_videos.items()},
        'profiles':{},'top_videos':[]}
    for k,p in profiles.items():
        result['profiles'][k]={'n':p['n'],'words':p['words'].most_common(50),
            'pairs':p['pairs'].most_common(40),'themes':p['themes'].most_common()}
    for vid,n in by_video.most_common(15):
        result['top_videos'].append({'video_id':vid,'title':lookup.get(vid,{}).get('title',''), 'n':n})
    result['top10_share']=round(sum(n for _,n in by_video.most_common(10))/max(1,quality['analysables'])*100,2)
    return result

HTML = '''<!doctype html><html lang="fr"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Commentaires YouTube — exploration du corpus</title>
<style>body{font:16px system-ui;margin:auto;max-width:1150px;padding:24px;background:#f4f6fa;color:#192b40}h1{font-size:32px}section,.card{background:white;padding:22px;border-radius:12px;margin:18px 0}nav a{margin-right:16px}a{color:#1855a5}button,select,input{padding:10px;margin:5px;border:1px solid #bbc7d6;border-radius:6px}table{width:100%;border-collapse:collapse}td,th{text-align:left;padding:8px;border-bottom:1px solid #e4e8ed}small,.muted{color:#53657a}.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(220px,1fr));gap:16px}.bar{display:block;background:#3779b5;height:7px;border-radius:3px;margin-top:5px}.notice{border-left:5px solid #d48b1f;padding:14px;background:#fff8e9}blockquote{margin:15px 0;padding:16px;background:#f4f6fa;white-space:pre-wrap;overflow-wrap:anywhere}.metric{font-size:30px;font-weight:bold}.tag{background:#e7effa;border-radius:4px;padding:3px 6px;font-size:12px}</style>
<h1>Ce qui revient dans les commentaires YouTube</h1><p>Paris · manifestation recherchée du 17 octobre 2026 · rapport exploratoire</p>
<p class="notice"><strong>Ce rapport ne mesure pas encore la polarisation.</strong> Il décrit le vocabulaire du corpus collecté. Les vidéos restent à valider; les commentaires YouTube ne représentent pas la population. Aucun score d’opinion, d’idéologie ou de radicalité n’est attribué aux personnes.</p>
<nav><a href="#qualite">Tri</a><a href="#frequences">Fréquences</a><a href="#themes">Thèmes</a><a href="#exemples">Exemples</a><a href="#methode">Méthode</a></nav>
<div id="metrics" class="grid"></div>
<section id="qualite"><h2>1. Qualité et premiers filtres</h2><div id="quality"></div><p>Les doublons sont des textes identiques après normalisation de la casse, des espaces et des entités HTML. Ce n’est pas une détection de bots. Un même texte peut provenir de plusieurs personnes. Les commentaires sans mot retenu incluent notamment emojis seuls, liens seuls et réponses très courtes; ils ne sont pas nécessairement inutiles pour une étude ultérieure.</p></section>
<section id="frequences"><h2>2. Mots et expressions récurrents</h2><label>Comparer : <select id="profile"><option value="tous">Tous les commentaires analysables</option><option value="uniques">Textes uniques analysables</option><option value="principaux">Commentaires principaux</option><option value="reponses">Réponses</option></select></label><p id="denominator"></p><div class="grid"><div><h3>Mots</h3><table id="words"></table></div><div><h3>Expressions de deux mots</h3><table id="pairs"></table></div></div><p class="muted">Nombre et pourcentage de commentaires contenant le terme au moins une fois. Les expressions conservent l’adjacence originale. Les formes fléchies restent distinctes.</p></section>
<section id="themes"><h2>3. Repérage thématique par mots-clés</h2><table id="themeTable"></table><p>Catégories définies à l’avance, non exclusives et non exhaustives. « Police » peut apparaître dans un commentaire favorable, hostile ou descriptif. Une mention n’est ni une position ni un sentiment. Les pourcentages ne s’additionnent pas à 100 %.</p></section>
<section><h2>4. Concentration du corpus par vidéo</h2><p id="concentration"></p><table id="videos"></table></section>
<section id="exemples"><h2>5. Échantillon de lecture et d’annotation</h2><p>Au plus un commentaire analysable par vidéo, sélectionné avec une graine fixe, puis 400 vidéos au maximum. Cet échantillon facilite la diversité de lecture; il n’est pas représentatif des commentaires. Extraits limités à 1 200 caractères; masquage automatique partiel des liens, emails, numéros et mentions. Les noms dans le texte peuvent subsister.</p><input id="search" placeholder="Rechercher dans les extraits" aria-label="Rechercher dans les extraits"><select id="themeFilter" aria-label="Filtrer par thème"><option value="">Tous les thèmes</option></select><p id="sampleCount"></p><div id="examples"></div><button id="more">Afficher 20 extraits supplémentaires</button></section>
<section id="methode"><h2>6. Méthode, limites et suite de l’étude</h2><p id="timestamp"></p><p>Analyse locale sans modèle externe : suppression des URLs pour le comptage, minuscules, accents normalisés, mots latins normalisés (a–z) d’au moins trois lettres; les écritures non latines ne sont pas analysées, liste explicite de mots vides français et anglais. Aucune lemmatisation ni détection automatique de langue. Les accents sont conservés dans les extraits. Comptage par présence dans chaque commentaire, et non par répétition d’un mot à l’intérieur du commentaire. Les réponses sont distinguées des commentaires principaux. Les thèmes sont des dictionnaires visibles dans analyze.py.</p><p>Le corpus est un résultat de recherche, pas un échantillon statistique. La couverture des chaînes, la visibilité algorithmique, les commentaires manquants et le poids des vidéos influencent les résultats. Un texte unique ne correspond pas nécessairement à une personne unique. Le masquage des extraits n’est pas une anonymisation garantie.</p><p>Pour étudier la polarisation : valider les vidéos et les sources; définir les objets du désaccord (mobilisation, maintien de l’ordre, violences, revendications); annoter un échantillon avec position, cible et incertitude; comparer les annotations de plusieurs personnes; puis seulement tester des mesures de désaccord et comparer les résultats par vidéo et par source. Les auteurs de commentaires n’ont pas été collectés : les réseaux d’interaction entre personnes ne sont pas reconstructibles avec ces données.</p><p>Fichiers joints : <a href="frequences.csv">fréquences CSV</a>, <a href="themes.csv">thèmes CSV</a>, <a href="annotation_400.csv">400 extraits à annoter</a>, <a href="synthese.json">synthèse JSON</a>. Le CSV d’annotation contient des champs manuels vides. Ce rapport fonctionne hors ligne et ne charge aucune ressource externe.</p></section>
<script id="payload" type="application/json">__DATA__</script><script>
const D=JSON.parse(document.getElementById('payload').textContent),$=id=>document.getElementById(id),fmt=n=>n.toLocaleString('fr-FR'),esc=s=>String(s??'').replace(/[&<>"']/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]));
function table(id,headers,rows){$(id).innerHTML='<thead><tr>'+headers.map(x=>'<th>'+esc(x)+'</th>').join('')+'</tr></thead><tbody>'+rows.map(r=>'<tr>'+r.map(x=>'<td>'+x+'</td>').join('')+'</tr>').join('')+'</tbody>'}
$('metrics').innerHTML=[[D.videos,'Vidéos candidates'],[D.total,'Commentaires et réponses'],[D.quality.analysables||0,'Commentaires analysables'],[D.profiles.uniques.n,'Textes uniques analysables']].map(([n,t])=>'<div class="card"><div class="metric">'+fmt(n)+'</div>'+esc(t)+'</div>').join('');
$('quality').innerHTML='<ul>'+[['vides','Textes vides'],['doublons_textuels','Occurrences répétées de textes normalisés'],['sans_mot_analyse','Sans mot retenu pour l’analyse'],['moins_de_3_mots_distincts','Analysables avec moins de trois mots distincts']].map(([k,label])=>'<li>'+esc(label)+' : '+fmt(D.quality[k]||0)+'</li>').join('')+'</ul><p>Les catégories de qualité peuvent se recouper. Aucun commentaire original n’est supprimé.</p>';
function render(){const p=D.profiles[$('profile').value];$('denominator').textContent=fmt(p.n)+' commentaires analysables dans cette vue.';for(const [key,id] of [['words','words'],['pairs','pairs']])table(id,['Terme','Commentaires','Part'],p[key].slice(0,30).map(([t,n])=>[esc(t),fmt(n),(100*n/Math.max(1,p.n)).toFixed(1)+' %<span class="bar" style="width:'+Math.round(100*n/Math.max(1,p.n))+'%"></span>']));table('themeTable',['Thème','Commentaires','Part de cette vue'],p.themes.map(([t,n])=>[esc(t),fmt(n),(100*n/Math.max(1,p.n)).toFixed(1)+' %']));}
$('profile').onchange=render;render();
$('concentration').textContent='Les dix vidéos les plus commentées regroupent '+D.top10_share+' % des commentaires analysables. '+fmt(D.videos_with_analyzable_comments)+' vidéos ont au moins un commentaire analysable. Le tableau ci-dessous présente les 15 premières.';
table('videos',['Vidéo','Commentaires analysables'],D.top_videos.map(v=>['<a target="_blank" rel="noopener noreferrer" href="https://www.youtube.com/watch?v='+encodeURIComponent(v.video_id)+'">'+esc(v.title)+'</a>',fmt(v.n)]));
[...new Set(D.sample.flatMap(x=>x.themes))].sort().forEach(t=>{let o=document.createElement('option');o.value=t;o.textContent=t;$('themeFilter').append(o)});
let limit=20;function examples(){let term=$('search').value.toLocaleLowerCase(),theme=$('themeFilter').value;let rows=D.sample.filter(r=>(!term||(r.text+' '+r.title).toLocaleLowerCase().includes(term))&&(!theme||r.themes.includes(theme)));$('sampleCount').textContent=rows.length+' extraits correspondent aux filtres; '+Math.min(limit,rows.length)+' affichés.';$('examples').innerHTML=rows.slice(0,limit).map(r=>'<article><p><a target="_blank" rel="noopener noreferrer" href="'+esc(r.url)+'">'+esc(r.title)+'</a> · '+esc(r.type)+'</p><blockquote>'+esc(r.text)+'</blockquote><p>'+r.themes.map(t=>'<span class="tag">'+esc(t)+'</span>').join(' ')+'</p></article>').join('');$('more').hidden=limit>=rows.length;}for(const id of ['search','themeFilter'])$(id).oninput=()=>{limit=20;examples()};$('more').onclick=()=>{limit+=20;examples()};examples();$('timestamp').textContent='Rapport généré le '+new Date(D.created_at).toLocaleString('fr-FR',{timeZone:'Europe/Paris'})+' (Europe/Paris).';
</script></html>'''

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input',type=Path)
    parser.add_argument('--output',type=Path,default=Path('reports'))
    args=parser.parse_args()
    source=args.input or next(p for p in [Path('data/paris_2026-10-17.json'),Path('data/paris_2026-10-17.xlsx')] if p.exists())
    videos,comments=load(source)
    result=analyze(videos,comments)
    result['source_file']=source.name
    result['source_sha256']=hashlib.sha256(source.read_bytes()).hexdigest()
    args.output.mkdir(parents=True,exist_ok=True)
    payload=json.dumps(result,ensure_ascii=False).replace('<','\\u003c').replace('>','\\u003e').replace('&','\\u0026')
    (args.output/'rapport.html').write_text(HTML.replace('__DATA__',payload),encoding='utf-8')
    (args.output/'synthese.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    csv_write(args.output/'frequences.csv',['vue','type','terme','commentaires','denominateur'],
        ([view,kind,term,n,p['n']] for view,p in result['profiles'].items() for kind in ['words','pairs'] for term,n in p[kind]))
    csv_write(args.output/'themes.csv',['vue','theme','commentaires','denominateur'],
        ([view,theme,n,p['n']] for view,p in result['profiles'].items() for theme,n in p['themes']))
    csv_write(args.output/'annotation_400.csv',['video_id','comment_id','video','texte_extrait','type','themes_mots_cles','video_pertinente','cible_position','position','incertitude','note'],
        ([r['video_id'],r['comment_id'],r['title'],r['text'],r['type'],' | '.join(r['themes']),'','','','',''] for r in result['sample']))
    print(json.dumps({'total':result['total'],'quality':result['quality'],'top_words':result['profiles']['tous']['words'][:15],'top10_share':result['top10_share']},ensure_ascii=False,indent=2))

if __name__=='__main__': main()
