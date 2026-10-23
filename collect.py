"""Collecte via YouTube Data API v3; aucune clé enregistrée sur disque."""
import argparse
import hashlib
from concurrent.futures import ThreadPoolExecutor
from threading import Lock, Event
import json
import os
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlencode
from urllib.request import urlopen
from urllib.error import HTTPError
from openpyxl import Workbook

QUERIES = ['manifestation Paris 17 octobre 2026', 'manifestations Paris 17/10/2026',
           'Paris protest October 17 2026', 'rassemblement Paris 17 octobre 2026']

class APIError(Exception):
    pass

class API:
    def __init__(self, key, cache=None):
        self.key = key
        self.cache = Path(cache) if cache else None
        if self.cache: self.cache.mkdir(parents=True, exist_ok=True)

    def call(self, resource, **params):
        cache_file = None
        if self.cache:
            fingerprint = hashlib.sha256(json.dumps([resource, params], sort_keys=True).encode()).hexdigest()
            cache_file = self.cache / (fingerprint + '.json')
            if cache_file.exists(): return json.loads(cache_file.read_text())
        url = 'https://www.googleapis.com/youtube/v3/' + resource + '?' + urlencode(dict(params, key=self.key))
        try:
            with urlopen(url, timeout=60) as response:
                result = json.load(response)
                if cache_file:
                    temporary = cache_file.with_suffix('.tmp')
                    temporary.write_text(json.dumps(result, ensure_ascii=False), encoding='utf-8')
                    temporary.replace(cache_file)
                return result
        except HTTPError as error:
            try:
                payload = json.load(error)
                reasons = [x.get('reason', 'unknown') for x in payload.get('error', {}).get('errors', [])]
            except (ValueError, KeyError):
                reasons = ['responseUnreadable']
            raise APIError(f'HTTP {error.code}: {", ".join(reasons)}') from None
        except OSError:
            raise APIError('Connexion API impossible; vérifier le réseau.') from None

    def pages(self, resource, **params):
        while True:
            result = self.call(resource, **params)
            yield from result.get('items', [])
            token = result.get('nextPageToken')
            if not token:
                break
            params['pageToken'] = token


def collect(api, checkpoint=None):
    videos, comments, errors = {}, {}, []
    for query in QUERIES:
        try:
            for item in api.pages('search', part='snippet', type='video', q=query,
                                  publishedAfter='2026-10-16T22:00:00Z',
                                  maxResults=50, order='relevance'):
                vid = item['id']['videoId']
                if vid not in videos:
                    videos[vid] = dict(video_id=vid, url='https://www.youtube.com/watch?v='+vid,
                        title=item['snippet']['title'], channel=item['snippet']['channelTitle'],
                        published_at=item['snippet']['publishedAt'], queries=[],
                        relevance='à vérifier manuellement', comments_status='non collectés')
                videos[vid]['queries'].append(query)
        except APIError as error:
            errors.append(dict(stage='search', target=query, error=str(error)))
    print(f'Recherche : {len(videos)} vidéos candidates', flush=True)
    def save():
        if checkpoint:
            checkpoint.parent.mkdir(parents=True, exist_ok=True)
            checkpoint.write_text(json.dumps(dict(videos=list(videos.values()), comments=list(comments.copy().values()), errors=errors), ensure_ascii=False, indent=2), encoding='utf-8')
    save()
    def add(item, vid, parent=''):
        s = item['snippet']
        comments[item['id']] = dict(comment_id=item['id'], video_id=vid, parent_id=parent,
            text=s.get('textOriginal', s.get('textDisplay', '')),
            published_at=s.get('publishedAt', ''), updated_at=s.get('updatedAt', ''),
            likes=s.get('likeCount', 0))
    completed = [0]
    lock = Lock()
    quota_stop = Event()
    def process(entry):
        vid, video = entry
        if quota_stop.is_set():
            video["comments_status"] = "non collectés — quota épuisé"
            return
        try:
            for thread in api.pages('commentThreads', part='snippet,replies', videoId=vid,
                                    textFormat='plainText', maxResults=100):
                top = thread['snippet']['topLevelComment']
                add(top, vid)
                replies = thread.get('replies', {}).get('comments', [])
                for reply in replies: add(reply, vid, top['id'])
                if thread['snippet'].get('totalReplyCount', 0) > len(replies):
                    for reply in api.pages('comments', part='snippet', parentId=top['id'],
                                           textFormat='plainText', maxResults=100):
                        add(reply, vid, top['id'])
            video['comments_status'] = 'pagination terminée pour les commentaires publics accessibles'
        except APIError as error:
            if 'quotaExceeded' in str(error) or 'dailyLimitExceeded' in str(error): quota_stop.set()
            video['comments_status'] = 'collecte incomplète ou commentaires indisponibles'
            errors.append(dict(stage='comments', target=vid, error=str(error)))
        with lock:
            completed[0] += 1
            if completed[0] % 10 == 0: save()
            print(f'Vidéos traitées : {completed[0]}/{len(videos)}, commentaires : {len(comments)}', flush=True)
    with ThreadPoolExecutor(max_workers=8) as pool:
        list(pool.map(process, videos.items()))
    save()
    return list(videos.values()), list(comments.copy().values()), errors


def export(path, videos, comments, errors, status):
    wb = Workbook()
    wb.remove(wb.active)
    metadata = [dict(field='statut', value=status),
        dict(field='collecte_utc', value=datetime.now(timezone.utc).isoformat()),
        dict(field='événement recherché', value='Paris, 17 octobre 2026 (Europe/Paris)'),
        dict(field='requêtes', value=' | '.join(QUERIES)),
        dict(field='limites', value='Résultats de recherche non exhaustifs; pertinence à vérifier. '
             'Commentaires publics uniquement, hors contenus supprimés, privés ou modérés. '
             'Les dates de publication ne prouvent pas la date des événements filmés.'),
        dict(field='données personnelles', value='Noms et identifiants des auteurs de commentaires omis; '
             'le texte peut encore contenir des données personnelles.'),
        dict(field='vidéos', value='Liens YouTube uniquement; aucun fichier vidéo téléchargé.')]
    tables = [('Métadonnées', metadata, ['field','value']),
              ('Vidéos', videos, ['video_id','url','title','channel','published_at','queries','relevance','comments_status']),
              ('Commentaires', comments, ['comment_id','video_id','parent_id','text','published_at','updated_at','likes']),
              ('Erreurs', errors, ['stage','target','error'])]
    for name, rows, keys in tables:
        # Répartition avant la limite Excel (1 048 576 lignes, en-tête compris).
        for offset in range(0, max(1, len(rows)), 1000000):
            ws = wb.create_sheet(name if offset == 0 else name+str(offset//1000000+1))
            ws.append(keys)
            for row_index, row in enumerate(rows[offset:offset+1000000], start=2):
                values = []
                for key in keys:
                    value = row.get(key, '')
                    if isinstance(value, list): value = ' | '.join(value)
                    if isinstance(value, str):
                        value = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', value)
                        if len(value) > 32767: raise ValueError('Texte trop long pour Excel')
                    values.append(value)
                ws.append(values)
                for column in range(1, len(keys)+1):
                    cell = ws.cell(row_index, column)
                    if isinstance(cell.value, str): cell.data_type = 's'
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--template', action='store_true', help='Créer un modèle explicitement vide')
    parser.add_argument('--output', type=Path, default=Path('data/paris_2026-10-17.xlsx'))
    args = parser.parse_args()
    if args.template:
        export(args.output, [], [], [], 'MODÈLE VIDE — aucune collecte effectuée')
        return 0
    key = os.environ.get('YOUTUBE_API_KEY')
    if not key:
        print('YOUTUBE_API_KEY manquante. Fournir la clé via les paramètres sécurisés, jamais dans Git.', file=sys.stderr)
        return 2
    videos, comments, errors = collect(API(key, '.youtube-cache'), args.output.with_suffix('.json'))
    # JSON conserve les données intégrales avant export Excel.
    args.output.with_suffix('.json').parent.mkdir(parents=True, exist_ok=True)
    args.output.with_suffix('.json').write_text(json.dumps(dict(videos=videos, comments=comments,
        errors=errors), ensure_ascii=False, indent=2), encoding='utf-8')
    export(args.output, videos, comments, errors, 'INCOMPLET — voir Erreurs' if errors else 'Collecte terminée; pertinence à vérifier, exhaustivité non garantie')
    print(f'{len(videos)} vidéos, {len(comments)} commentaires, {len(errors)} erreurs')
    return 1 if errors else 0

if __name__ == '__main__':
    sys.exit(main())
