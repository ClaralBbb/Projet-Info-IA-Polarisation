import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from collect import API, APIError, collect, export

class Tests(unittest.TestCase):
    def test_pagination(self):
        api = API('fake')
        calls = []
        def call(resource, **params):
            calls.append(params)
            return {'items': [2]} if params.get('pageToken') else {'items': [1], 'nextPageToken': 'next'}
        api.call = call
        self.assertEqual(list(api.pages('search')), [1, 2])
        self.assertEqual(calls[1]['pageToken'], 'next')

    def test_collect_and_export(self):
        class Fake:
            def pages(self, resource, **params):
                if resource == 'search':
                    yield {'id': {'videoId': 'v'}, 'snippet': {'title': '=unsafe', 'channelTitle': 'c', 'publishedAt': '2026-10-18'}}
                elif resource == 'commentThreads':
                    yield {'snippet': {'topLevelComment': {'id': 'top', 'snippet': {'textOriginal': '=1+1'}}, 'totalReplyCount': 1}}
                else:
                    yield {'id': 'reply', 'snippet': {'textOriginal': 'réponse'}}
        videos, comments, errors = collect(Fake())
        self.assertEqual(len(videos), 1)
        self.assertEqual(len(comments), 2)
        self.assertEqual(comments[1]['parent_id'], 'top')
        self.assertFalse(errors)
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)/'result.xlsx'
            export(path, videos, comments, errors, 'test')
            wb = load_workbook(path)
            self.assertEqual(wb['Commentaires'].max_row, 3)
            self.assertEqual(wb['Commentaires']['D2'].data_type, 's')
            self.assertEqual(wb['Commentaires']['D2'].value, '=1+1')

    def test_error_not_silently_successful(self):
        class Broken:
            def pages(self, resource, **params):
                raise APIError('HTTP 403: quotaExceeded')
        videos, comments, errors = collect(Broken())
        self.assertFalse(videos)
        self.assertEqual(len(errors), 4)

    def test_cache_reuses_response_without_saving_key(self):
        from unittest.mock import patch
        from io import BytesIO
        with tempfile.TemporaryDirectory() as folder:
            api = API('secret-test-key', folder)
            with patch('collect.urlopen', return_value=BytesIO(b'{"items": []}')) as request:
                self.assertEqual(api.call('search', q='Paris'), {'items': []})
                self.assertEqual(api.call('search', q='Paris'), {'items': []})
                self.assertEqual(request.call_count, 1)
            for path in Path(folder).iterdir():
                self.assertNotIn('secret-test-key', path.read_text())
                self.assertNotIn('secret-test-key', path.name)

    def test_multiple_videos_keep_errors_and_comments_separate(self):
        class Fake:
            def pages(self, resource, **params):
                if resource == 'search':
                    for vid in ['good', 'disabled']:
                        yield {'id': {'videoId': vid}, 'snippet': {'title': vid, 'channelTitle': 'c', 'publishedAt': '2026-10-17'}}
                elif params['videoId'] == 'disabled':
                    raise APIError('HTTP 403: commentsDisabled')
                else:
                    yield {'snippet': {'topLevelComment': {'id': 'comment', 'snippet': {'textOriginal': 'ok'}}, 'totalReplyCount': 0}}
        videos, comments, errors = collect(Fake())
        self.assertEqual(len(videos), 2)
        self.assertEqual(comments[0]['video_id'], 'good')
        self.assertEqual(errors[0]['target'], 'disabled')
        self.assertIn('incomplète', videos[1]['comments_status'])

if __name__ == '__main__': unittest.main()
